from contextlib import contextmanager
import json

from openpyxl import load_workbook
import pytest

from tiktok_categorizer import cli
from tiktok_categorizer.captions import fill_missing_captions, lookup_caption
from tiktok_categorizer.categories import parse_categories
from tiktok_categorizer.collector import CollectionStopped
from tiktok_categorizer.models import SavedVideo, load_video_list, read_video_list, save_video_list

URL = "https://www.tiktok.com/@janestyle/video/7301234567890123456"
OEMBED = {"title": "fit check #ootd", "author_name": "Jane", "author_url": "https://www.tiktok.com/@janestyle"}


def test_oembed_fills_in_caption_and_creator():
    video = SavedVideo(url=URL)
    seen = []

    def http_get(url):
        seen.append(url)
        return 200, json.dumps(OEMBED).encode()

    assert lookup_caption(video, http_get)
    assert (video.caption, video.creator) == ("fit check #ootd", "@janestyle")
    assert seen == ["https://www.tiktok.com/oembed?url=https%3A%2F%2Fwww.tiktok.com%2F%40janestyle%2Fvideo%2F7301234567890123456"]


def test_oembed_retries_rate_limits_but_not_missing_videos():
    replies = [(429, b""), (200, json.dumps(OEMBED).encode())]
    sleeps = []
    assert lookup_caption(SavedVideo(url=URL), lambda _: replies.pop(0), sleep=sleeps.append)
    assert sleeps == [1]

    calls = []
    assert not lookup_caption(SavedVideo(url=URL), lambda u: calls.append(u) or (404, b""), sleep=sleeps.append)
    assert len(calls) == 1


def test_only_videos_with_nothing_known_are_looked_up():
    videos = [SavedVideo(url=URL), SavedVideo(url=URL, creator="@x"), SavedVideo(url=URL, caption="hi")]
    calls = []

    def http_get(url):
        calls.append(url)
        return 200, json.dumps(OEMBED).encode()

    assert fill_missing_captions(videos, http_get) == 1
    assert len(calls) == 1


def _saved(tmp_path, name="list.json", *, username="me", complete=True):
    path = tmp_path / name
    save_video_list(path, [
        SavedVideo(url=URL, caption="fit check #ootd", creator="@janestyle", posted_at="2023-11-14"),
        SavedVideo(url=URL.replace("3456", "3457"), caption="donate to the #gofundme", creator="@helper"),
        SavedVideo(url=URL.replace("3456", "3458"), caption="", creator="@quiet"),
    ], username, complete=complete)
    return path


def test_video_lists_remember_whether_they_are_finished(tmp_path):
    finished = read_video_list(_saved(tmp_path, "a.json"))
    unfinished = read_video_list(_saved(tmp_path, "b.json", complete=False))
    assert (finished.complete, unfinished.complete, unfinished.username) == (True, False, "me")
    assert len(unfinished.videos) == 3
    older = tmp_path / "older.json"  # lists written before this flag existed count as finished
    older.write_text(json.dumps({"username": "me", "videos": [{"url": URL}]}))
    assert read_video_list(older).complete
    assert not list(tmp_path.glob("*.tmp"))


def test_resorting_a_saved_list_writes_the_spreadsheet(tmp_path, capsys):
    out = tmp_path / "videos.xlsx"
    assert cli.main(["--from-list", str(_saved(tmp_path)), "-o", str(out), "--no-open"]) == 0

    book = load_workbook(out)
    assert {"Fashion", "Fundraisers & Causes", "Uncategorized"} <= set(book.sheetnames)
    assert "Fashion: 1" in capsys.readouterr().err
    assert not (tmp_path / "videos.json").exists()  # the input list is not rewritten


def test_from_list_alone_uses_the_list_next_to_the_spreadsheet(tmp_path):
    _saved(tmp_path, "videos.json")
    out = tmp_path / "videos.xlsx"
    assert cli.main(["--from-list", "-o", str(out), "--no-open"]) == 0
    assert out.exists()


class FakeTikTok:
    """Stands in for the browser part of the command."""

    def __init__(self, monkeypatch, result=None, stop=None, stop_with=None):
        self.calls = []
        self.result, self.stop, self.stop_with = result or [], stop, stop_with

        @contextmanager
        def fake_browser(profile_dir, **kwargs):
            yield "context", "page"

        monkeypatch.setattr(cli, "open_tiktok_browser", fake_browser)
        monkeypatch.setattr(cli, "wait_for_login", lambda context, page, **kwargs: None)
        monkeypatch.setattr(cli, "detect_username", lambda page: "me")
        monkeypatch.setattr(cli, "collect_favorites", self.collect)

    def collect(self, page, username, *, limit, known, on_progress, should_stop, log):
        self.calls.append({"username": username, "limit": limit, "known": len(known)})
        found = known + [v for v in self.result if v.key not in {k.key for k in known}]
        on_progress(found[:2])
        if self.stop:
            raise CollectionStopped(self.stop, found if self.stop_with is None else self.stop_with)
        return found[:limit] if limit else found


def _three(tmp_path):
    return load_video_list(_saved(tmp_path, "source.json"))


def _progress(tmp_path, user="me"):
    return tmp_path / f"mine.{user}.in-progress.json"


def test_a_full_run_saves_a_finished_list(tmp_path, monkeypatch):
    tiktok = FakeTikTok(monkeypatch, result=_three(tmp_path))
    out = tmp_path / "mine.xlsx"
    assert cli.main(["-o", str(out), "--no-open"]) == 0
    saved = read_video_list(tmp_path / "mine.json")
    assert (len(saved.videos), saved.username, saved.complete) == (3, "me", True)
    assert tiktok.calls == [{"username": "me", "limit": None, "known": 0}]
    assert out.exists() and not _progress(tmp_path).exists()


def test_a_stopped_run_keeps_progress_and_writes_a_so_far_spreadsheet(tmp_path, monkeypatch, capsys):
    FakeTikTok(monkeypatch, result=_three(tmp_path), stop="TikTok's page kept crashing")
    out = tmp_path / "mine.xlsx"
    assert cli.main(["-o", str(out), "--no-open"]) == 1
    progress = read_video_list(_progress(tmp_path))
    assert (len(progress.videos), progress.complete) == (3, False)
    assert not (tmp_path / "mine.json").exists() and not out.exists()
    assert load_workbook(tmp_path / "mine (so far).xlsx")["All Videos"].max_row == 4
    assert "Run the same command again" in capsys.readouterr().err


def test_a_stopped_refresh_never_replaces_the_finished_list_or_spreadsheet(tmp_path, monkeypatch):
    finished = _saved(tmp_path, "mine.json")
    out = tmp_path / "mine.xlsx"
    out.write_bytes(b"the full spreadsheet")
    before = finished.read_text()
    FakeTikTok(monkeypatch, result=_three(tmp_path)[:1], stop="Stopped.")
    assert cli.main(["-o", str(out), "--no-open"]) == 1
    assert finished.read_text() == before and out.read_bytes() == b"the full spreadsheet"
    assert len(read_video_list(_progress(tmp_path)).videos) == 1


def test_a_stop_with_nothing_found_keeps_earlier_progress(tmp_path, monkeypatch):
    progress = _saved(tmp_path, "mine.me.in-progress.json", complete=False)
    before = progress.read_text()
    FakeTikTok(monkeypatch, result=[], stop="Stopped.", stop_with=[])
    assert cli.main(["-o", str(tmp_path / "mine.xlsx"), "--no-open", "--start-over"]) == 1
    assert progress.read_text() == before


def test_the_next_run_continues_an_unfinished_one_and_tidies_up(tmp_path, monkeypatch):
    _saved(tmp_path, "mine.me.in-progress.json", complete=False)
    (tmp_path / "mine (so far).xlsx").write_bytes(b"old partial")
    tiktok = FakeTikTok(monkeypatch, result=[SavedVideo(url=URL.replace("3456", "3999"), caption="new")])
    assert cli.main(["-o", str(tmp_path / "mine.xlsx"), "--no-open"]) == 0
    assert tiktok.calls[0]["known"] == 3
    assert len(read_video_list(tmp_path / "mine.json").videos) == 4
    assert not _progress(tmp_path).exists() and not (tmp_path / "mine (so far).xlsx").exists()


def test_other_accounts_progress_and_start_over_are_not_continued(tmp_path, monkeypatch):
    tiktok = FakeTikTok(monkeypatch, result=[])
    out = str(tmp_path / "mine.xlsx")
    _saved(tmp_path, "mine.someone_else.in-progress.json", username="someone_else", complete=False)
    cli.main(["-o", out, "--no-open"])
    _saved(tmp_path, "mine.me.in-progress.json", complete=False)
    cli.main(["-o", out, "--no-open", "--start-over"])
    assert [call["known"] for call in tiktok.calls] == [0, 0]
    assert len(read_video_list(_progress(tmp_path, "someone_else")).videos) == 3


def test_a_limited_test_run_leaves_saved_results_alone(tmp_path, monkeypatch):
    progress = _saved(tmp_path, "mine.me.in-progress.json", complete=False)
    before = progress.read_text()
    tiktok = FakeTikTok(monkeypatch, result=_three(tmp_path))
    assert cli.main(["-o", str(tmp_path / "mine.xlsx"), "--limit", "1", "--no-open"]) == 0
    assert tiktok.calls[0]["known"] == 0
    assert progress.read_text() == before
    assert (tmp_path / "mine (test).xlsx").exists()
    assert not (tmp_path / "mine.xlsx").exists() and not (tmp_path / "mine.json").exists()


def test_a_stopped_test_run_says_nothing_was_saved(tmp_path, monkeypatch, capsys):
    FakeTikTok(monkeypatch, result=_three(tmp_path), stop="Stopped.")
    assert cli.main(["-o", str(tmp_path / "mine.xlsx"), "--limit", "5", "--no-open"]) == 1
    err = capsys.readouterr().err
    assert "test run" in err and "progress is saved" not in err
    assert not _progress(tmp_path).exists()


def test_from_list_falls_back_to_unfinished_progress(tmp_path):
    _saved(tmp_path, "mine.me.in-progress.json", complete=False)
    assert cli.main(["--from-list", "-o", str(tmp_path / "mine.xlsx"), "--no-open"]) == 0
    assert (tmp_path / "mine.xlsx").exists()


def test_control_c_while_sorting_says_how_to_finish(tmp_path, monkeypatch, capsys):
    def interrupted(*args, **kwargs):
        raise KeyboardInterrupt

    monkeypatch.setattr(cli, "fill_missing_captions", interrupted)
    assert cli.main(["--from-list", str(_saved(tmp_path)), "-o", str(tmp_path / "x.xlsx"), "--ai"]) == 130
    assert "tiktok-categorize --from-list --ai" in capsys.readouterr().err


def test_first_control_c_asks_to_stop_and_the_second_quits(monkeypatch):
    exits = []
    monkeypatch.setattr(cli.os, "_exit", exits.append)
    stop = cli._GracefulStop()
    assert not stop.requested()
    stop._handle(2, None)
    assert stop.requested() and not exits
    stop._handle(2, None)
    assert exits == [130]


def test_ai_mode_without_the_package_or_key_falls_back(tmp_path, monkeypatch, capsys):
    from tiktok_categorizer import ai

    def unavailable(*args, **kwargs):
        raise ai.AIUnavailable("no key")

    monkeypatch.setattr(ai, "AIClassifier", unavailable)
    out = tmp_path / "videos.xlsx"
    assert cli.main(["--from-list", str(_saved(tmp_path)), "-o", str(out), "--ai", "--no-open"]) == 0
    assert "keyword rules instead" in capsys.readouterr().err
    assert out.exists()


def test_default_categories_print_as_a_usable_file(capsys):
    assert cli.main(["--print-default-categories"]) == 0
    assert parse_categories(json.loads(capsys.readouterr().out))


def test_results_go_to_documents_by_default():
    assert cli.DEFAULT_OUTPUT.parent.name == "TikTok-Saved-Videos"
    assert cli.DEFAULT_OUTPUT.parent.parent.name == "Documents"


def test_bad_inputs_exit_with_a_message(tmp_path, capsys):
    bad = tmp_path / "cats.json"
    bad.write_text("[]")
    out = str(tmp_path / "out.xlsx")
    assert cli.main(["--from-list", str(_saved(tmp_path)), "--categories", str(bad), "-o", out]) == 1
    assert cli.main(["--from-list", str(tmp_path / "missing.json"), "-o", out]) == 1
    assert cli.main(["--from-list", "-o", str(tmp_path / "nothing" / "yet.xlsx")]) == 1
    assert cli.main(["--limit", "0"]) == 2
    err = capsys.readouterr().err
    assert "Error:" in err and "--limit" in err
